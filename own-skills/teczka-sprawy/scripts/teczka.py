#!/usr/bin/env python3
"""Dokumenty sprawy w teczce legal/<sprawa>/ (port teczka_kancelaria.py z abogada).

  00_BRIEF.docx         z analiza/_brief.md (albo _brief.md w korzeniu teczki)
  01_SPIS_TRESCI.docx   z analiza/_dokumenty.md (spis dokumentów sprawy w ../dokumentacja) oraz plików w teczce
                        poza materiałem roboczym (analiza/, drafty/, _robocze/)
  02_PODSUMOWANIE.docx  z analiza/_podsumowanie.md (albo _podsumowanie.md)

Brief i podsumowanie redaguje się w plikach .md; skrypt tylko je składa (pandoc). Uruchom po każdej zmianie.

  python3 scripts/teczka.py --sprawa <katalog sprawy>
  python3 scripts/teczka.py --wszystkie
"""
from __future__ import annotations

import argparse
import os
import datetime as dt
import re
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(os.environ.get("LEGAL_WORKSPACE", os.getcwd())).resolve()
ROBOCZE = {"analiza", "drafty", "_robocze", "@eaDir"}
RODZAJ = {".pdf": "PDF", ".docx": "Word", ".doc": "Word", ".jpg": "zdjęcie", ".jpeg": "zdjęcie", ".png": "obraz",
          ".heic": "zdjęcie", ".m4a": "nagranie audio", ".mp3": "nagranie audio", ".mp4": "nagranie wideo",
          ".mov": "nagranie wideo", ".eml": "e-mail", ".html": "HTML", ".md": "tekst", ".xlsx": "arkusz",
          ".json": "dane", ".txt": "tekst"}


def sprawy() -> list[Path]:
    """Teczka sprawy = katalog w korzeniu legal z analiza/ albo _brief.md."""
    return sorted(d for d in ROOT.iterdir()
                  if d.is_dir() and not d.name.startswith(("_", ".", "@", "#", "00_"))
                  and ((d / "analiza").is_dir() or (d / "_brief.md").exists()))


def zrodlo(sprawa: Path, nazwa: str) -> Path | None:
    for p in (sprawa / "analiza" / nazwa, sprawa / nazwa):
        if p.exists():
            return p
    return None


def data_z_nazwy(n: str) -> str:
    m = re.search(r"(20\d{2})[-_.]?(\d{2})[-_.]?(\d{2})", n)
    if m and 1 <= int(m.group(2)) <= 12 and 1 <= int(m.group(3)) <= 31:
        return f"{m.group(3)}.{m.group(2)}.{m.group(1)}"
    return ""


def karta(n: str) -> str:
    m = re.search(r"(?:^|[_ -])k\.?\s?(\d+(?:[-–]k?\d+)?)", n, re.I)
    return f"k. {m.group(1).replace('k', '')}" if m else ""


def strony(p: Path) -> str:
    if p.suffix.lower() != ".pdf" or not p.exists():
        return ""
    try:
        import pymupdf
        with pymupdf.open(p) as d:
            return str(d.page_count)
    except Exception:  # noqa: BLE001 - brak PyMuPDF albo uszkodzony PDF nie blokuje spisu
        return ""


def dokumenty_z_rejestru(sprawa: Path) -> list[Path]:
    """Pliki wymienione w _dokumenty.md (sekcje `## \\`<katalog>\\``, wiersze `| \\`plik\\` | ... |`)."""
    rej = zrodlo(sprawa, "_dokumenty.md")
    if not rej:
        return []
    out, katalog = [], None
    for linia in rej.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^##\s+`([^`]+)`", linia)
        if m:
            katalog = (ROOT / m.group(1)).resolve() if not m.group(1).startswith("/") else Path(m.group(1))
            continue
        m = re.match(r"^\|\s*`([^`]+)`\s*\|", linia)
        if m and katalog is not None:
            out.append(katalog / m.group(1))
    return out


def tabela(doc, pliki: list[Path], baza: Path | None) -> int:
    from docx.shared import Pt
    t = doc.add_table(rows=1, cols=6)
    t.style = "Light Grid Accent 1"
    for c, h in zip(t.rows[0].cells, ("Lp.", "Data", "Dokument", "Rodzaj", "Str.", "Karta")):
        c.text = h
    klucz = lambda x: (data_z_nazwy(x.name)[-4:] + data_z_nazwy(x.name)[3:5] + data_z_nazwy(x.name)[:2], x.name)
    for lp, p in enumerate(sorted(pliki, key=klucz), 1):
        nazwa = str(p.relative_to(baza)) if baza and p.is_relative_to(baza) else p.name
        vals = (str(lp), data_z_nazwy(p.name), nazwa, RODZAJ.get(p.suffix.lower(), p.suffix.lstrip(".")),
                strony(p), karta(p.name))
        for c, v in zip(t.add_row().cells, vals):
            c.text = v
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8)
    return len(pliki)


def spis(sprawa: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor
    doc = Document()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(f"Projekt - do weryfikacji przed przekazaniem | {dt.date.today().strftime('%d.%m.%Y')}")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    doc.add_heading("Spis treści teczki", level=0)
    doc.add_paragraph(sprawa.name.replace("_", " "))
    razem = 0
    rejestr = dokumenty_z_rejestru(sprawa)
    if rejestr:
        grupy: dict[Path, list[Path]] = {}
        for x in rejestr:
            grupy.setdefault(x.parent, []).append(x)
        for kat, pliki in grupy.items():
            try:
                etykieta = str(kat.relative_to(ROOT.parent))
            except ValueError:
                etykieta = str(kat)
            doc.add_heading(etykieta, level=1)
            razem += tabela(doc, pliki, kat)
    wlasne = [x for x in sorted(sprawa.rglob("*"))
              if x.is_file() and not x.name.startswith((".", "~$", "_"))
              and not (x.parent == sprawa and x.name[:2] in ("00", "01", "02"))
              and x.relative_to(sprawa).parts[0] not in ROBOCZE]
    if wlasne:
        doc.add_heading("Dokumenty w teczce", level=1)
        razem += tabela(doc, wlasne, sprawa)
    doc.add_heading("Materiał roboczy (poza teczką)", level=1)
    for k, opis in (("analiza", "analizy i rejestry sprawy"), ("drafty", "projekty pism"),
                    ("_robocze", "wyniki skilli kontroli i pliki jednorazowe")):
        d = sprawa / k
        n = sum(1 for x in d.rglob("*") if x.is_file() and "@eaDir" not in x.parts) if d.exists() else 0
        doc.add_paragraph(f"{k}/ - {n} plików: {opis}", style="List Bullet")
    doc.save(sprawa / "01_SPIS_TRESCI.docx")
    print(f"  01_SPIS_TRESCI.docx: {razem} dokumentów")


def z_markdown(sprawa: Path, nazwa: str, cel: str, tytul: str) -> None:
    src = zrodlo(sprawa, nazwa)
    if not src:
        print(f"  pominięto {cel}: brak {nazwa}")
        return
    if not shutil.which("pandoc"):
        print(f"  ! brak pandoc - nie mogę złożyć {cel}")
        return
    meta = (f"---\ntitle: \"{tytul}\"\nsubtitle: \"{sprawa.name.replace('_', ' ')} - projekt do weryfikacji, "
            f"{dt.date.today().strftime('%d.%m.%Y')}\"\nlang: pl-PL\n---\n\n")
    subprocess.run(["pandoc", "-f", "markdown", "-o", str(sprawa / cel)],
                   input=(meta + src.read_text(encoding="utf-8")).encode(), check=True)
    print(f"  {cel}: z {src.relative_to(sprawa)}")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--sprawa", help="nazwa katalogu sprawy w legal/")
    g.add_argument("--wszystkie", action="store_true")
    a = ap.parse_args()
    try:
        import docx  # noqa: F401
    except ImportError:
        print("brak python-docx: pip install --break-system-packages python-docx", file=sys.stderr)
        return 2
    lista = sprawy() if a.wszystkie else [ROOT / a.sprawa]
    for sprawa in lista:
        if not sprawa.is_dir():
            print(f"brak katalogu sprawy: {sprawa}", file=sys.stderr)
            return 1
        print(f"== {sprawa.name}")
        z_markdown(sprawa, "_brief.md", "00_BRIEF.docx", "Brief sprawy")
        spis(sprawa)
        z_markdown(sprawa, "_podsumowanie.md", "02_PODSUMOWANIE.docx", "Podsumowanie sprawy")
    return 0


if __name__ == "__main__":
    sys.exit(main())
