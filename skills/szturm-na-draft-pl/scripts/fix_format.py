#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1"]
# ///
"""Poprawki formatu .docx, ktorych adeu nie sledzi (E4/E7 szturmu).

Wykrywa (zawsze) i naprawia (tylko z --apply):
  NAGLOWEK_TEKST  akapit tekstu ciaglego w stylu naglowka (Heading/Naglowek N): > --slowa slow
                  albo konczy sie kropka. Krotka etykieta z dwukropkiem ("Zalaczniki:") zostaje.
                  Naprawa: styl tekstu glownego dokumentu.
  PODKRESLENIE    podkreslony run w tekscie ciaglym (poza naglowkami i hiperlaczami).
                  Naprawa: zdjecie podkreslenia.
  NAGLOWEK_PUSTY  akapit w stylu naglowka bez tekstu (pusta pozycja w nawigacji i spisie tresci).
                  Naprawa: styl tekstu glownego. Akapitu nie usuwamy - moze trzymac podzial strony.

Zmiany z --apply NIE sa zmianami sledzonymi Worda. Dlatego skrypt:
  - nigdy nie nadpisuje wejscia (wymaga -o),
  - wypisuje raport JSON kazdej zmiany (akapit, stary styl, tekst), ktory idzie do raportu
    szturmu w sekcji "Zmiany nieslledzone".

Kolejnosc w E7: najpierw adeu apply (zmiany sledzone), potem ten skrypt na pliku REDLINE
albo na wersji czystej - nigdy odwrotnie, bo adeu dopasowuje tekst i styl po naszej zmianie
moglby rozjechac kotwice.

Uruchamianie (python-docx nie jest w systemowym Pythonie - nie instaluj go globalnie):
  uv run <skille>/szturm-na-draft-pl/scripts/fix_format.py pismo.docx                 # raport
  uv run <skille>/szturm-na-draft-pl/scripts/fix_format.py pismo.docx --apply -o pismo_fmt.docx
  opcje: --slowa 12  --bez-podkreslen  --bez-naglowkow  --json

Exit: 0 = brak ustalen albo zastosowano, 1 = sa ustalenia (tryb raportu), 2 = blad.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from collections import Counter
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    print("BLAD: brak python-docx. Uruchom przez: uv run fix_format.py ...", file=sys.stderr)
    sys.exit(2)

RE_HEADING = re.compile(r"^(heading|nagłówek|naglowek|title|tytuł)\s*\d*$", re.I)


def czy_naglowek(par) -> bool:
    st = par.style
    while st is not None:
        if RE_HEADING.match((st.name or "").strip()):
            return True
        st = st.base_style
    ppr = par._p.pPr
    return ppr is not None and ppr.find(qn("w:outlineLvl")) is not None


def w_hiperlaczu(run) -> bool:
    return run._r.getparent().tag == qn("w:hyperlink")


def styl_tekstu(doc) -> object:
    """Najczestszy styl akapitow niebedacych naglowkami - to on jest 'tekstem glownym'."""
    licz = Counter(p.style.name for p in doc.paragraphs if p.text.strip() and not czy_naglowek(p))
    if not licz:
        return doc.styles["Normal"]
    return doc.styles[licz.most_common(1)[0][0]]


def skanuj(doc, slowa: int, podkr: bool, nagl: bool) -> list[dict]:
    wyniki = []
    for i, p in enumerate(doc.paragraphs, 1):
        tekst = p.text.strip()
        if czy_naglowek(p):
            if not nagl:
                continue
            if not tekst:
                wyniki.append({"typ": "NAGLOWEK_PUSTY", "akapit": i, "styl": p.style.name, "tekst": ""})
            elif len(tekst.split()) > slowa or tekst.endswith("."):
                wyniki.append({"typ": "NAGLOWEK_TEKST", "akapit": i, "styl": p.style.name,
                               "slow": len(tekst.split()), "tekst": tekst[:120]})
            continue
        if podkr:
            for r in p.runs:
                if r.font.underline and r.text.strip() and not w_hiperlaczu(r):
                    wyniki.append({"typ": "PODKRESLENIE", "akapit": i, "styl": p.style.name,
                                   "tekst": r.text.strip()[:120]})
    return wyniki


def zastosuj(doc, wyniki: list[dict]) -> list[dict]:
    cel = styl_tekstu(doc)
    pars = doc.paragraphs
    zmiany = []
    for w in wyniki:
        p = pars[w["akapit"] - 1]
        if w["typ"] in ("NAGLOWEK_TEKST", "NAGLOWEK_PUSTY"):
            stary = p.style.name
            p.style = cel
            ppr = p._p.pPr
            if ppr is not None and ppr.find(qn("w:outlineLvl")) is not None:
                ppr.remove(ppr.find(qn("w:outlineLvl")))
            zmiany.append({**w, "zmiana": f"styl '{stary}' -> '{cel.name}'"})
        elif w["typ"] == "PODKRESLENIE":
            for r in p.runs:
                if r.font.underline and r.text.strip()[:120] == w["tekst"] and not w_hiperlaczu(r):
                    r.font.underline = None
                    zmiany.append({**w, "zmiana": "zdjete podkreslenie"})
                    break
    return zmiany


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx", type=Path)
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("-o", "--output", type=Path)
    ap.add_argument("--slowa", type=int, default=12, help="prog slow dla tekstu w stylu naglowka")
    ap.add_argument("--bez-podkreslen", action="store_true")
    ap.add_argument("--bez-naglowkow", action="store_true")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()
    try:
        doc = Document(str(a.docx))
    except Exception as e:  # noqa: BLE001 - python-docx rzuca rozne wyjatki dla zlego pliku
        print(f"BLAD: {a.docx}: {e}", file=sys.stderr)
        return 2
    wyniki = skanuj(doc, a.slowa, not a.bez_podkreslen, not a.bez_naglowkow)
    if a.apply:
        if not a.output or a.output.resolve() == a.docx.resolve():
            print("BLAD: --apply wymaga -o innego niz plik wejsciowy (zmiany nie sa sledzone)", file=sys.stderr)
            return 2
        # najpierw style, potem ponowny skan podkreslen: akapit zdjety z naglowka staje sie
        # tekstem ciaglym i jego podkreslenia dopiero wtedy sa ustaleniem
        zmiany = zastosuj(doc, [w for w in wyniki if w["typ"] != "PODKRESLENIE"])
        if not a.bez_podkreslen:
            zmiany += zastosuj(doc, [w for w in skanuj(doc, a.slowa, True, False)])
        doc.save(str(a.output))
        raport = {"plik": str(a.docx), "wynik": str(a.output), "zmiany_niesledzone": zmiany}
    else:
        raport = {"plik": str(a.docx), "ustalenia": wyniki}
    if a.json:
        print(json.dumps(raport, ensure_ascii=False, indent=2))
    else:
        licz = Counter(w["typ"] for w in wyniki)
        print(f"{a.docx.name}: " + (", ".join(f"{k} {v}" for k, v in sorted(licz.items())) or "brak ustalen"))
        for w in raport.get("zmiany_niesledzone", raport.get("ustalenia", [])):
            print(f"  akapit {w['akapit']:>4} {w['typ']:<15} [{w['styl']}] {w['tekst'][:80]}"
                  + (f"  => {w['zmiana']}" if "zmiana" in w else ""))
    if a.apply:
        return 0
    return 1 if wyniki else 0


if __name__ == "__main__":
    sys.exit(main())
