#!/usr/bin/env python3
# /// script
# requires-python = ">=3.10"
# dependencies = ["python-docx>=1.1"]
# ///
"""Syntetyczne pismo do testu dymnego szturmu. Wszystkie osoby, sygnatury i fakty sa fikcyjne.

Kolejnosc akapitow jest czescia kontraktu testu: kotwice w fixture_ustalenia.json to linie
w <nazwa>.clean.md z adeu (naglowek pliku + akapity rozdzielone pustymi liniami).
"""
import sys

from docx import Document

doc = Document()
doc.add_heading("PISMO TESTOWE - STANOWISKO", level=1)
doc.add_paragraph("Sąd powołał się na wyrok z dnia 15 marca 2024 r., sygn. II K 000/23, w którym przyjęto definicję zamiaru.")
doc.add_paragraph("Obwiniony widział zapis z tej kamery z dźwiękiem w dniu 1 sierpnia 2025 r.")
doc.add_paragraph("Ścieżka dźwiękowa zawiera sygnał o poziomie szczytowym zmierzonym w pliku.")
doc.add_paragraph("Z 1000 próbek większość ma wartość zero.")
doc.add_paragraph("Obrona opiera się na braku dowodu skutku.")
doc.add_paragraph("Nadto zapis nie zawiera obrazu otwierającego sekwencję.")
doc.add_heading("To jest zwykły akapit tekstu ciągłego omyłkowo sformatowany jako nagłówek i zakończony kropką.", level=1)
p = doc.add_paragraph("Wniosek dowodowy dotyczy ")
p.add_run("nagrania porównawczego").underline = True
p.add_run(" z tej samej kamery.")
doc.add_paragraph("Nagranie zawiera wyraźny sygnał mowy w końcowym odcinku.")
doc.add_paragraph("Protokół oględzin nie odnotowuje zniszczenia kamery.")
doc.add_paragraph("Co więcej, wniosek jest zasadny.")
doc.save(sys.argv[1])
