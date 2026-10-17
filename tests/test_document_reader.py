import importlib.util
import json
from pathlib import Path
import sys
from unittest.mock import patch

import pymupdf
import pytest

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT/'plugins/grounding-prawo/skills/odczyt-dokumentow'
spec = importlib.util.spec_from_file_location('reader', SKILL/'scripts/read_documents.py')
reader = importlib.util.module_from_spec(spec)
spec.loader.exec_module(reader)
from the_ocr.models import OCRDocument, OCRPage, OCRBlock
from the_ocr.adapters import datalab_document, polygon_document


def pdf(path, text='Sample document with evidence and source text.'):
    with pymupdf.open() as doc:
        page = doc.new_page();page.insert_text((50,50), text);doc.save(path)


def test_native_pdf_preserved_without_ocr(tmp_path):
    p=tmp_path/'input.pdf';pdf(p);before=p.read_bytes()
    with patch.object(reader, 'cloud_ocr', side_effect=AssertionError('OCR must not run')):
        result=reader.convert(p,tmp_path/'out')
    assert p.read_bytes()==before
    contract=json.loads(Path(result['contract']).read_text())
    assert contract['source']['engine']=='pymupdf'
    assert contract['blocks'][0]['bbox'] is not None
    assert '## Page 1' in Path(result['markdown']).read_text()


@pytest.mark.parametrize('provider',['datalab','azure-di','google-vision'])
def test_provider_contract_and_original_page_geometry(tmp_path,provider):
    p=tmp_path/'input.pdf';pdf(p)
    if provider=='datalab':
        doc=datalab_document({'page_count':1,'json':{'children':[{'bbox':[0,0,600,800],
            'children':[{'block_type':'Text','bbox':[60,80,300,160],'html':'<p>Evidence text</p>'}]}]}},p)
    else:
        doc=polygon_document([{'width':600,'height':800,'lines':[{'text':'Evidence text',
            'polygon':[60,80,300,80,300,160,60,160],
            'words':[{'text':'Evidence','polygon':[60,80,300,80,300,160,60,160],'confidence':.96}]}]}],p,provider=provider)
    with patch.object(reader,'cloud_ocr',return_value=doc) as cloud:
        result=reader.convert(p,tmp_path/'out',provider,True)
    cloud.assert_called_once_with(p,provider)
    c=json.loads(Path(result['contract']).read_text())
    assert not reader.validate(c)
    assert c['source']['engine']==provider
    assert c['blocks'][0]['bbox']==[.1,.1,.5,.2]
    assert c['blocks'][0]['page']==1
    if provider=='datalab': assert c['blocks'][0]['confidence'] is None


def test_duplicate_and_cache(tmp_path):
    src=tmp_path/'inputs';src.mkdir()
    (src/'a.txt').write_text('One document')
    (src/'b.txt').write_text('One document')
    output=tmp_path/'out'
    result=reader.read_documents(src,output)
    assert len(result['duplicates']['identical_files'])==1
    assert result['documents'][1]['cached']
    with patch.object(reader,'extract',side_effect=AssertionError('cache missed')):
        assert reader.read_documents(src,output)['documents'][0]['cached']
    (src/'a.txt').write_text('Changed document')
    assert not reader.read_documents(src,output)['documents'][0]['cached']


def test_unsupported_is_failure_and_empty_never_success(tmp_path):
    src=tmp_path/'inputs';src.mkdir()
    (src/'empty.txt').write_text('')
    (src/'unknown.xyz').write_text('test')
    results=reader.read_documents(src,tmp_path/'out')['documents']
    assert all(r['status']=='failed' for r in results)


def test_docx_and_html_no_invented_pagination(tmp_path):
    from docx import Document
    p=tmp_path/'input.docx';doc=Document();doc.add_paragraph('Treść dokumentu');doc.save(p)
    h=tmp_path/'input.html';h.write_text('<html><body><h1>Tytuł</h1><p>Treść dokumentu</p></body></html>')
    for path in (p,h):
        result=reader.convert(path,tmp_path/'out')
        contract=json.loads(Path(result['contract']).read_text())
        assert contract['source']['pages'] is None
        assert contract['blocks'][0]['bbox'] is None
        assert 'Treść dokumentu' in Path(result['markdown']).read_text()
        assert '## Page' not in Path(result['markdown']).read_text()


def test_empty_page_report(tmp_path):
    p=tmp_path/'input.pdf';pdf(p)
    doc=OCRDocument('input.pdf','hash','azure-di',(
        OCRPage(1,600,800,(OCRBlock('b1',1,'text','Text','Text',(.1,.1,.5,.2),.99),)),
        OCRPage(2,600,800,())))
    with patch.object(reader,'cloud_ocr',return_value=doc):
        result=reader.convert(p,tmp_path/'out',force_ocr=True)
    assert result['quality']['empty_or_unreadable_pages']==[2]
    assert result['status']=='degraded'


def test_markdown_layout_from_provider_survives(tmp_path):
    p=tmp_path/'input.pdf';pdf(p)
    table='| A | B |\n|---|---|\n| 1 | 2 |'
    doc=OCRDocument('input.pdf','hash','datalab',(
        OCRPage(1,600,800,(OCRBlock('b1',1,'table','A B 1 2',table,(.1,.1,.5,.2),None),)),))
    with patch.object(reader,'cloud_ocr',return_value=doc):
        result=reader.convert(p,tmp_path/'out',force_ocr=True)
    assert table in Path(result['markdown']).read_text()


def test_missing_build_input_does_not_erase_existing_package(tmp_path):
    import subprocess
    (tmp_path/'scripts').mkdir();(tmp_path/'selection').mkdir()
    (tmp_path/'scripts/build.py').write_text((ROOT/'scripts/build.py').read_text())
    (tmp_path/'selection/selection.yaml').write_text('keep:\n  - {name: missing, from: own}\ndrop: []\n')
    marker=tmp_path/'plugins/grounding-prawo/skills/existing/SKILL.md'
    marker.parent.mkdir(parents=True);marker.write_text('Existing package')
    run=subprocess.run([sys.executable,str(tmp_path/'scripts/build.py')],capture_output=True,text=True)
    assert run.returncode!=0
    assert marker.read_text()=='Existing package'
